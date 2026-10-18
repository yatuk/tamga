"""Text extraction from attached files (app.extract).

Every file is made here, with the libraries that read them or by hand
(PDF), so nothing is downloaded and no fixture is checked in. The tests of
a format are skipped when its library is not installed.
"""

from __future__ import annotations

import asyncio
import io
import struct
import time
import zipfile

import pytest

from app.extract import (
    STATUS_ENCRYPTED,
    STATUS_EXTRACT_FAILED,
    STATUS_OCR_UNAVAILABLE,
    STATUS_OK,
    STATUS_TOO_LARGE,
    STATUS_UNSUPPORTED_TYPE,
    Extractor,
    Limits,
    extract_bytes,
)
from app.extract import readers, sniff

TCKN = "10000000146"  # the checksum-valid test identity number used across the repo


# ── File builders ─────────────────────────────────────────────────────────


def _pdf(objects: list[bytes]) -> bytes:
    """A PDF from its objects (1-based, object 1 the catalog), with a
    correct cross-reference table."""
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % n + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def _stream(data: bytes, extra: bytes = b"") -> bytes:
    return b"<< /Length %d %s >>\nstream\n" % (len(data), extra) + data + b"\nendstream"


def text_pdf(pages: list[str]) -> bytes:
    """A PDF with one line of text on each page."""
    n = len(pages)
    # 1 catalog, 2 pages, 3 font, then a page and its content per page.
    kids = b" ".join(b"%d 0 R" % (4 + 2 * i) for i in range(n))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % n,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        content = b"BT /F1 12 Tf 72 712 Td (" + text.encode("latin-1") + b") Tj ET"
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i)
        )
        objects.append(_stream(content))
    return _pdf(objects)


def image_pdf() -> bytes:
    """A one-page PDF whose page is an image and nothing else: a scan."""
    pixels = bytes([0, 255, 255, 0])  # 2×2 grey
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /XObject << /Im1 4 0 R >> >> /Contents 5 0 R >>",
        _stream(pixels, b"/Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceGray /BitsPerComponent 8"),
        _stream(b"q 100 0 0 100 72 600 cm /Im1 Do Q"),
    ]
    return _pdf(objects)


def encrypted_pdf(user_password: str) -> bytes:
    pypdf = pytest.importorskip("pypdf")
    writer = pypdf.PdfWriter(clone_from=io.BytesIO(text_pdf([f"secret {TCKN}"])))
    writer.encrypt(user_password=user_password, owner_password="owner", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_file() -> bytes:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph(f"Customer identity number {TCKN}")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "IBAN"
    table.rows[0].cells[1].text = "TR00 0000 0000 0000 0000 0000 00"
    document.sections[0].header.paragraphs[0].text = "Header line"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_file() -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["name", "id"])
    ws.append(["Test Person", TCKN])
    other = wb.create_sheet("Second")
    other["A1"] = "on the second sheet"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pptx_file(slides: int = 1) -> bytes:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i + 1}"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text_frame.text = f"body {i + 1} {TCKN}"
        slide.notes_slide.notes_text_frame.text = f"speaker notes {i + 1}"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def png_file(width: int = 4, height: int = 4) -> bytes:
    pil = pytest.importorskip("PIL.Image")
    out = io.BytesIO()
    pil.new("L", (width, height), color=128).save(out, format="PNG")
    return out.getvalue()


def set_zip_flag(data: bytes, bit: int) -> bytes:
    """Sets a general-purpose flag bit on every local and central header."""
    buf = bytearray(data)
    for sig, off in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        i = buf.find(sig)
        while i >= 0:
            flags = struct.unpack_from("<H", buf, i + off)[0]
            struct.pack_into("<H", buf, i + off, flags | bit)
            i = buf.find(sig, i + 4)
    return bytes(buf)


def declare_size(data: bytes, size: int) -> bytes:
    """Makes every central directory entry claim to unpack to size bytes."""
    buf = bytearray(data)
    i = buf.find(b"PK\x01\x02")
    while i >= 0:
        struct.pack_into("<I", buf, i + 24, size)
        i = buf.find(b"PK\x01\x02", i + 4)
    return bytes(buf)


def fake_ocr(image: bytes) -> str:
    return f"read by ocr {TCKN}"


# ── Formats ───────────────────────────────────────────────────────────────


def test_pdf_text_per_page():
    pytest.importorskip("pypdf")
    res = extract_bytes(text_pdf(["first page", f"second {TCKN}"]))
    assert res.status == STATUS_OK, res
    assert res.detected_type == "pdf"
    assert [(p.kind, p.index) for p in res.parts] == [("page", 1), ("page", 2)]
    assert TCKN in res.parts[1].text
    assert not res.truncated


def test_pdf_falls_back_to_pdfminer(monkeypatch):
    pypdf = pytest.importorskip("pypdf")
    pytest.importorskip("pdfminer")
    monkeypatch.setattr(pypdf.PageObject, "extract_text", lambda self, *a, **kw: "")
    res = extract_bytes(text_pdf([f"via pdfminer {TCKN}"]))
    assert res.status == STATUS_OK, res
    assert TCKN in res.parts[0].text


def test_pdf_encrypted_with_a_password():
    res = extract_bytes(encrypted_pdf("secret"))
    assert res.status == STATUS_ENCRYPTED
    assert res.parts == []


def test_pdf_encrypted_with_an_empty_user_password_is_read():
    res = extract_bytes(encrypted_pdf(""))
    assert res.status == STATUS_OK, res
    assert TCKN in res.parts[0].text


def test_scanned_pdf_goes_to_ocr():
    pytest.importorskip("pypdf")
    pytest.importorskip("PIL")
    assert extract_bytes(image_pdf()).status == STATUS_OCR_UNAVAILABLE
    res = extract_bytes(image_pdf(), ocr=fake_ocr)
    assert res.status == STATUS_OK, res
    assert res.parts[0].ocr and TCKN in res.parts[0].text


def test_docx_paragraphs_tables_and_headers():
    res = extract_bytes(docx_file())
    assert res.status == STATUS_OK, res
    assert res.detected_type == "docx"
    text = res.parts[0].text
    assert TCKN in text and "TR00" in text and "Header line" in text
    # The table cell's paragraph is read once.
    assert text.count("IBAN") == 1


def test_xlsx_sheets_with_names():
    res = extract_bytes(xlsx_file())
    assert res.status == STATUS_OK, res
    assert [(p.kind, p.name) for p in res.parts] == [("sheet", "Customers"), ("sheet", "Second")]
    assert TCKN in res.parts[0].text


def test_xlsx_formula_reads_its_cached_value():
    xlsxwriter = pytest.importorskip("xlsxwriter")
    pytest.importorskip("openpyxl")
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out)
    ws = wb.add_worksheet("F")
    ws.write_formula("A1", '=CONCAT("100000","00146")', None, TCKN)
    wb.close()
    res = extract_bytes(out.getvalue())
    assert res.status == STATUS_OK, res
    assert TCKN in res.parts[0].text
    assert "CONCAT" not in res.parts[0].text


def test_pptx_slides_and_notes():
    res = extract_bytes(pptx_file(2))
    assert res.status == STATUS_OK, res
    assert [p.kind for p in res.parts] == ["slide", "slide"]
    assert "Slide 1" in res.parts[0].text and TCKN in res.parts[0].text
    assert "speaker notes 2" in res.parts[1].text


def test_image_needs_ocr():
    data = png_file()
    assert extract_bytes(data).status == STATUS_OCR_UNAVAILABLE
    res = extract_bytes(data, ocr=fake_ocr)
    assert res.status == STATUS_OK and res.parts[0].kind == "image" and res.parts[0].ocr


def test_image_with_too_many_pixels_is_not_decoded(monkeypatch):
    monkeypatch.setattr("app.extract.limits.MAX_IMAGE_PIXELS", 10)
    assert extract_bytes(png_file(8, 8), ocr=fake_ocr).status == STATUS_TOO_LARGE


# ── Refusals ──────────────────────────────────────────────────────────────


def test_encrypted_zip_entry_is_refused_before_reading():
    data = set_zip_flag(docx_file(), 0x1)
    res = extract_bytes(data)
    assert res.status == STATUS_ENCRYPTED and res.detected_type == "docx"


def test_password_protected_office_file():
    cfb = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504)
    cfb += "EncryptionInfo".encode("utf-16-le") + bytes(100) + "EncryptedPackage".encode("utf-16-le")
    assert extract_bytes(cfb).status == STATUS_ENCRYPTED
    legacy = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504) + "WordDocument".encode("utf-16-le")
    assert extract_bytes(legacy).status == STATUS_UNSUPPORTED_TYPE


def test_zip_bomb_by_declared_size():
    data = declare_size(docx_file(), 0xF0000000)
    assert extract_bytes(data).status == STATUS_TOO_LARGE


def test_zip_bomb_by_ratio():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", b"\0" * (24 << 20))
    data = out.getvalue()
    assert len(data) < (1 << 20)
    assert extract_bytes(data).status == STATUS_TOO_LARGE


def test_unsupported_types():
    plain_zip = io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as zf:
        zf.writestr("a.txt", "hello")
    for data in (plain_zip.getvalue(), b"\x1f\x8b\x08\x00", b"just some bytes", b""):
        assert extract_bytes(data).status == STATUS_UNSUPPORTED_TYPE, data[:8]


def test_over_max_bytes():
    pytest.importorskip("pypdf")
    data = text_pdf(["x"])
    assert extract_bytes(data, Limits.of(max_bytes=len(data) - 1)).status == STATUS_TOO_LARGE


def test_damaged_files_are_a_status_not_an_exception():
    pytest.importorskip("pypdf")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("[Content_Types].xml", "not xml <<<")
        zf.writestr("word/document.xml", "not xml either")
        zf.writestr("xl/workbook.xml", "nor this")
    cases = [
        out.getvalue(),
        b"%PDF-1.4\n" + bytes(range(256)) * 20,
        text_pdf(["cut"])[:60],
        docx_file()[:500],
    ]
    for data in cases:
        res = extract_bytes(data)
        assert res.status in (STATUS_EXTRACT_FAILED, STATUS_UNSUPPORTED_TYPE, STATUS_OK), res
        assert res.status != STATUS_OK or res.parts == [] or isinstance(res.parts[0].text, str)


# ── Limits ────────────────────────────────────────────────────────────────


def test_max_pages_truncates():
    pytest.importorskip("pypdf")
    res = extract_bytes(text_pdf(["one", "two", "three"]), Limits.of(max_pages=2))
    assert res.status == STATUS_OK
    assert [p.index for p in res.parts] == [1, 2]
    assert res.truncated
    res = extract_bytes(pptx_file(3), Limits.of(max_pages=1))
    assert len(res.parts) == 1 and res.truncated


def test_max_chars_truncates():
    res = extract_bytes(pptx_file(3), Limits.of(max_chars=12))
    assert res.truncated
    assert sum(len(p.text) for p in res.parts) <= 12


def test_limits_are_held_to_the_hard_bounds():
    lim = Limits.of(max_bytes=1 << 40, max_pages=10**6, max_chars=10**9, timeout_ms=10**7)
    assert lim.max_bytes == 32 << 20 and lim.max_pages == 2000
    assert lim.max_chars == 4_000_000 and lim.timeout_s == 60.0
    assert Limits.of() == Limits()


def test_deadline_between_parts():
    res = extract_bytes(pptx_file(2), deadline=time.monotonic() - 1)
    assert res.status == STATUS_EXTRACT_FAILED


# ── The worker pool ───────────────────────────────────────────────────────


def test_extractor_reads_and_answers_timeouts(monkeypatch):
    pytest.importorskip("pypdf")
    ex = Extractor(workers=1, ocr_factory=lambda: None)
    try:
        res = asyncio.run(ex.extract(text_pdf([f"x {TCKN}"]), Limits.of()))
        assert res.status == STATUS_OK and TCKN in res.parts[0].text

        def stuck(data, budget, ocr):
            time.sleep(1.5)
            return budget.result("pdf")

        monkeypatch.setattr(readers, "read_pdf", stuck)

        async def two():
            started = time.monotonic()
            first = asyncio.ensure_future(ex.extract(text_pdf(["y"]), Limits.of(timeout_ms=200)))
            await asyncio.sleep(0.05)
            # The one worker is held by the stuck reader: refused at once.
            second = await ex.extract(text_pdf(["z"]), Limits.of(timeout_ms=200))
            return await first, second, time.monotonic() - started

        first, second, took = asyncio.run(two())
        assert first.status == STATUS_EXTRACT_FAILED
        assert second.status == STATUS_EXTRACT_FAILED
        assert took < 1.4
    finally:
        ex.shutdown()


def test_extractor_ocr_engine(monkeypatch):
    class Engine:
        def available(self):
            return True

        def read(self, image, lang):
            assert lang == "tur"
            return type("R", (), {"text": f"engine text {TCKN}"})()

    ex = Extractor(workers=1, ocr_factory=Engine)
    try:
        res = asyncio.run(ex.extract(png_file(), Limits.of(), ocr=True, ocr_lang="tur"))
        assert res.status == STATUS_OK and TCKN in res.parts[0].text
        # Not asked for OCR: not read.
        assert asyncio.run(ex.extract(png_file(), Limits.of())).status == STATUS_OCR_UNAVAILABLE
    finally:
        ex.shutdown()


def test_no_ocr_module_means_ocr_unavailable():
    from app.extract.service import load_ocr_engine

    engine = load_ocr_engine()
    if engine is None:
        ex = Extractor(workers=1)
        try:
            res = asyncio.run(ex.extract(png_file(), Limits.of(), ocr=True))
            assert res.status == STATUS_OCR_UNAVAILABLE
        finally:
            ex.shutdown()


def test_sniff_does_not_trust_the_name():
    assert sniff.sniff(b"%PDF-1.4 ...").type == sniff.TYPE_PDF
    assert sniff.sniff(b"\x89PNG\r\n\x1a\n....").type == sniff.TYPE_IMAGE
