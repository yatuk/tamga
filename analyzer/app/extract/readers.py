"""Readers: the text of a PDF, DOCX, XLSX or PPTX file, or of an image.

Each reader takes the file's bytes and a Budget, which holds the limits of
the request: it stops at max_pages parts and max_chars characters (and says
the text was truncated), and raises DeadlineExceeded between parts once the
deadline has passed.
"""

from __future__ import annotations

import io
import logging
import time
from typing import Any, Callable

from app.extract import limits as L
from app.extract.result import (
    STATUS_ENCRYPTED,
    STATUS_EXTRACT_FAILED,
    STATUS_OCR_UNAVAILABLE,
    STATUS_OK,
    STATUS_TOO_LARGE,
    Part,
    Result,
)

logger = logging.getLogger("tamga.extract")


class DeadlineExceeded(Exception):
    """The extraction ran past its deadline."""


class Budget:
    """What is left of a request's limits as the parts are read."""

    def __init__(self, limits: L.Limits, deadline: float) -> None:
        self.limits = limits
        self.deadline = deadline
        self.chars_left = limits.max_chars
        self.truncated = False
        self.parts: list[Part] = []

    def check_time(self) -> None:
        if time.monotonic() > self.deadline:
            raise DeadlineExceeded()

    def page_allowed(self, index: int) -> bool:
        """Whether the part at index (0-based) may be read."""
        if index >= self.limits.max_pages:
            self.truncated = True
            return False
        return True

    @property
    def full(self) -> bool:
        return self.chars_left <= 0

    def add(self, kind: str, index: int, text: str, name: str = "", ocr: bool = False) -> None:
        """Keeps the text of one part, cut to what the budget has left."""
        text = (text or "").strip()
        if not text:
            return
        if self.chars_left <= 0:
            self.truncated = True
            return
        if len(text) > self.chars_left:
            text = text[: self.chars_left]
            self.truncated = True
        self.chars_left -= len(text)
        self.parts.append(Part(kind=kind, index=index, text=text, name=name, ocr=ocr))

    def result(self, detected: str) -> Result:
        return Result(status=STATUS_OK, parts=self.parts, truncated=self.truncated, detected_type=detected)


# ── OCR ───────────────────────────────────────────────────────────────────


OcrRead = Callable[[bytes], str]


def ocr_text(value: Any) -> str:
    """The text an OCR engine returned: a string, or an object with .text."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    text = getattr(value, "text", None)
    return text if isinstance(text, str) else ""


def image_too_large(data: bytes) -> bool:
    """Whether the image has more pixels than is decoded. The header is read,
    not the pixels. Without Pillow the check is left to the OCR engine."""
    try:
        from PIL import Image  # python-pptx depends on Pillow
    except ImportError:  # pragma: no cover - Pillow comes with python-pptx
        return False
    try:
        with Image.open(io.BytesIO(data)) as img:
            width, height = img.size
    except Exception:  # noqa: BLE001 - an unreadable header is the engine's to report
        return False
    return width * height > L.MAX_IMAGE_PIXELS


def read_image(data: bytes, budget: Budget, ocr: OcrRead | None) -> Result:
    if ocr is None:
        return Result(status=STATUS_OCR_UNAVAILABLE, detected_type="image")
    if image_too_large(data):
        return Result(status=STATUS_TOO_LARGE, detected_type="image")
    budget.add("image", 1, ocr(data), ocr=True)
    return budget.result("image")


# ── PDF ───────────────────────────────────────────────────────────────────


def read_pdf(data: bytes, budget: Budget, ocr: OcrRead | None) -> Result:
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError, PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except PdfReadError:
        return Result(status=STATUS_EXTRACT_FAILED, detected_type="pdf")
    if reader.is_encrypted:
        # Many PDFs are encrypted with an empty user password, to set
        # permissions: anyone can open them, and so can this. One that needs
        # a password, or a cipher this install cannot decrypt, is encrypted
        # as far as anyone reading it here is concerned.
        try:
            if not reader.decrypt(""):
                return Result(status=STATUS_ENCRYPTED, detected_type="pdf")
        except Exception:  # noqa: BLE001 - DependencyError, NotImplementedError, …
            return Result(status=STATUS_ENCRYPTED, detected_type="pdf")

    try:
        pages = reader.pages
        count = len(pages)
    except FileNotDecryptedError:
        return Result(status=STATUS_ENCRYPTED, detected_type="pdf")

    for i in range(count):
        if not budget.page_allowed(i) or budget.full:
            break
        budget.check_time()
        try:
            text = pages[i].extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page is not the file
            logger.warning("pdf page unreadable: %s", type(exc).__name__)
            continue
        budget.add("page", i + 1, text)
    if budget.parts:
        return budget.result("pdf")

    # pypdf found no text: pdfminer reads some layouts it does not.
    budget.check_time()
    _read_pdf_pdfminer(data, budget)
    if budget.parts:
        return budget.result("pdf")

    # No text layer: a scan. Its pages' images are read by OCR.
    images: list[tuple[int, bytes]] = []
    for i in range(min(count, budget.limits.max_pages)):
        budget.check_time()
        try:
            for img in pages[i].images:
                images.append((i + 1, img.data))
        except Exception as exc:  # noqa: BLE001
            logger.warning("pdf page images unreadable: %s", type(exc).__name__)
    if not images:
        # Nothing in it to read: a blank document is not an unreadable one.
        return budget.result("pdf")
    if ocr is None:
        return Result(status=STATUS_OCR_UNAVAILABLE, detected_type="pdf")
    for page, img in images:
        if budget.full:
            budget.truncated = True
            break
        budget.check_time()
        if image_too_large(img):
            budget.truncated = True
            continue
        budget.add("page", page, ocr(img), ocr=True)
    return budget.result("pdf")


def _read_pdf_pdfminer(data: bytes, budget: Budget) -> None:
    try:
        from pdfminer.high_level import extract_text
    except ImportError:  # pragma: no cover - a requirement
        return
    try:
        text = extract_text(io.BytesIO(data), maxpages=budget.limits.max_pages, password="")
    except Exception as exc:  # noqa: BLE001
        logger.warning("pdfminer could not read the pdf: %s", type(exc).__name__)
        return
    # pdfminer ends every page with a form feed.
    for i, page in enumerate(text.split("\f")):
        budget.add("page", i + 1, page)


# ── DOCX ──────────────────────────────────────────────────────────────────


def _paragraph_lines(root: Any) -> list[str]:
    """The text of every paragraph under root, tables and text boxes
    included, each paragraph once."""
    from docx.oxml.ns import qn

    p_tag, t_tag, tab_tag, br_tag = qn("w:p"), qn("w:t"), qn("w:tab"), qn("w:br")
    lines: list[str] = []
    for p in root.iter(p_tag):
        pieces: list[str] = []
        for el in p.iter(t_tag, tab_tag, br_tag):
            # A run in a text box inside this paragraph belongs to the text
            # box's own paragraph, which the outer loop reaches as well.
            owner = el.getparent()
            while owner is not None and owner.tag != p_tag:
                owner = owner.getparent()
            if owner is not p:
                continue
            if el.tag == t_tag:
                pieces.append(el.text or "")
            elif el.tag == tab_tag:
                pieces.append("\t")
            else:
                pieces.append("\n")
        line = "".join(pieces).strip()
        if line:
            lines.append(line)
    return lines


def read_docx(data: bytes, budget: Budget) -> Result:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = _paragraph_lines(document.element.body)
    for section in document.sections:
        budget.check_time()
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            try:
                if part is not None and not part.is_linked_to_previous:
                    lines.extend(_paragraph_lines(part._element))
            except Exception:  # noqa: BLE001 - a missing part is not an error
                continue
    budget.add("document", 1, "\n".join(lines))
    return budget.result("docx")


# ── XLSX ──────────────────────────────────────────────────────────────────


def read_xlsx(data: bytes, budget: Budget) -> Result:
    import openpyxl

    # data_only: a formula's value as the file last saved it, which is what
    # a reader of the sheet sees, not the formula.
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for i, ws in enumerate(wb.worksheets):
            if not budget.page_allowed(i) or budget.full:
                break
            rows: list[str] = []
            used = 0
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n % 256 == 0:
                    budget.check_time()
                cells = ["" if v is None else str(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells:
                    continue
                line = "\t".join(cells)
                rows.append(line)
                used += len(line) + 1
                if used > budget.chars_left:
                    budget.truncated = True
                    break
            budget.add("sheet", i + 1, "\n".join(rows), name=ws.title)
    finally:
        wb.close()
    return budget.result("xlsx")


# ── PPTX ──────────────────────────────────────────────────────────────────


def _shape_texts(shapes: Any, out: list[str], depth: int = 0) -> None:
    if depth > 16:
        return
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6 or hasattr(shape, "shapes"):  # a group
            try:
                _shape_texts(shape.shapes, out, depth + 1)
            except Exception:  # noqa: BLE001
                pass
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text
            if text.strip():
                out.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text for cell in row.cells]
                if any(c.strip() for c in cells):
                    out.append("\t".join(cells))


def read_pptx(data: bytes, budget: Budget) -> Result:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    for i, slide in enumerate(presentation.slides):
        if not budget.page_allowed(i) or budget.full:
            break
        budget.check_time()
        texts: list[str] = []
        _shape_texts(slide.shapes, texts)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                texts.append(notes.text)
        budget.add("slide", i + 1, "\n".join(texts))
    return budget.result("pptx")
